import os
from datetime import datetime
from docx import Document


def export_to_txt(content, file_name):

    os.makedirs("exports", exist_ok=True)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

    path = f"exports/{file_name}_{timestamp}.txt"

    with open(path, "w", encoding="utf-8") as file:
        file.write(content)

    return path


def export_to_docx(content, file_name):

    os.makedirs("exports", exist_ok=True)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")

    path = f"exports/{file_name}_{timestamp}.docx"

    document = Document()

    document.add_heading(file_name.replace("_", " ").title(), level=1)

    document.add_paragraph(content)

    document.save(path)

    return path